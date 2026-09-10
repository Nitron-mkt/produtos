#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Kit flexivel do Nitron Mob PDV: 3 paineis, 6 ripas, 4 conectores -> 45 modulos padrao
(3 footprints x 3 larguras x 5 pilhas). Gera dados/46-nitron-mob-kit-flex.xlsx, tudo em formula."""
import importlib.util, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
RAIZ=pathlib.Path(__file__).resolve().parent.parent
spec=importlib.util.spec_from_file_location('cad',RAIZ/'analise'/'pdv-caderno.py')
cad=importlib.util.module_from_spec(spec); spec.loader.exec_module(cad)
OUT=RAIZ/'dados'/'46-nitron-mob-kit-flex.xlsx'
F='Arial'; fb=Font(name=F,size=10); fB=Font(name=F,size=10,bold=True); fin=Font(name=F,size=10,color='0000FF')
fh=Font(name=F,size=10,bold=True,color='FFFFFF'); ft=Font(name=F,size=14,bold=True); fn=Font(name=F,size=9,italic=True,color='555555')
HEAD=PatternFill('solid',fgColor='1EA7AC'); YEL=PatternFill('solid',fgColor='FFFF00'); GREY=PatternFill('solid',fgColor='F2F2F2')
thin=Side(style='thin',color='BFBFBF'); box=Border(left=thin,right=thin,top=thin,bottom=thin)
C=Alignment(horizontal='center',vertical='center'); WR=Alignment(wrap_text=True,vertical='top')
BRL='"R$" #,##0.00'; KG='0.0'; MM='#,##0'
def header(ws,row,cols):
    for i,c in enumerate(cols):
        x=ws.cell(row,1+i,c); x.font=fh; x.fill=HEAD; x.alignment=C; x.border=box
def widths(ws,w):
    for i,x in enumerate(w,1): ws.column_dimensions[L(i)].width=x
GMM,DENS,PANT,RSKG=cad.GMM,cad.DENS,cad.PANT,cad.RSKG
ripa_kg=lambda mm: mm*GMM/1000; pan_kg=lambda l,c: l*c*PANT*DENS/1000
wb=Workbook()

# ============================================================ Parametros
P=wb.active; P.title='Parametros'
P['A1']='Cotas do sistema — amarelo = medido na malha STL, ainda sem trena'; P['A1'].font=fB
header(P,4,['Parâmetro','Símbolo','Valor (mm)'])
par=[('Encaixe','ENC',40.60,1),('Trizeta · comprimento','NOX',61.61,1),('Cruzeta · comprimento','NOXC',101.30,1),
     ('Trizeta · profundidade','NOY',83.23,1),('Trizeta · vertical','NOZ',73.08,1),('Peça L · sobra por lado','NOL',21.92,1),
     ('Pé exposto','PE',19.40,1),('Espessura do painel','PANT',15,0),('Ripa consome','CONSOME','=2*C5',0)]
for i,(a,b,v,chk) in enumerate(par):
    r=5+i; P.cell(r,1,a).font=fb; P.cell(r,2,b).font=fB
    c=P.cell(r,3,v); c.font=fb if isinstance(v,str) else fin
    if chk: c.fill=YEL
    for k in range(1,4): P.cell(r,k).border=box
widths(P,[28,10,12])
N_=dict(ENC='Parametros!$C$5',NOX='Parametros!$C$6',NOXC='Parametros!$C$7',NOY='Parametros!$C$8',NOZ='Parametros!$C$9',
        NOL='Parametros!$C$10',PE='Parametros!$C$11',PANT='Parametros!$C$12',CONS='Parametros!$C$13')

# ============================================================ Kit
K=wb.create_sheet('Kit')
K['A1']='Nitron Mob PDV — kit flexível'; K['A1'].font=ft
K['A2']='13 peças: 3 painéis, 6 ripas (todas PI existentes), 4 conectores. Custo unitário em amarelo é entrada; a madeira usa R$ %.2f/kg do caderno.'%RSKG; K['A2'].font=fn
header(K,4,['Peça','Código','Medida (mm)','Massa unit. (kg)','Custo unit. (R$)','Papel'])
pecas=[
 ('Painel 300×450','novo · pinus 15','300 × 450',pan_kg(300,450),pan_kg(300,450)*RSKG,'prateleira rasa curta — ponta, canto, caixa–pilar'),
 ('Painel 300×754','novo · pinus 15','300 × 754',pan_kg(300,754),pan_kg(300,754)*RSKG,'prateleira rasa padrão — 82% dos SKUs cabem'),
 ('Painel 460×754','novo · pinus 15','460 × 754',pan_kg(460,754),pan_kg(460,754)*RSKG,'prateleira funda — lixeira, cesto, caixa grande, e a parede mais estável'),
 ('Ripa largura 287','BLA-03-AC','287',ripa_kg(287),ripa_kg(287)*RSKG,'profundidade 372'),
 ('Ripa largura 415','PSC-02','415',ripa_kg(415),ripa_kg(415)*RSKG,'profundidade 500 (a mesma PI serve de vão 450)'),
 ('Ripa comprimento 415','PSC-02','415',ripa_kg(415),ripa_kg(415)*RSKG,'vão de 450'),
 ('Ripa comprimento 717','PSC-04','717',ripa_kg(717),ripa_kg(717)*RSKG,'vão de 754'),
 ('Ripa vertical 270','BAL-02-AC','270',ripa_kg(270),ripa_kg(270)*RSKG,'baia de 247 livres'),
 ('Ripa vertical 513','PSA-05','513',ripa_kg(513),ripa_kg(513)*RSKG,'baia de 490 livres — e a coroa da arara'),
 ('Pé','sobra de ripa','60',ripa_kg(60),ripa_kg(60)*RSKG,''),
 ('Trizeta','PP · molde existe','—',cad.M_TZ/1000,cad.C_TZ,'canto de cada nível'),
 ('Cruzeta','PP · 4 injetadas','—',cad.M_CZ/1000,cad.C_CZ,'só quando dois vãos dividem um poste (N ≥ 2)'),
 ('Peça L','PP · molde existe','—',cad.M_L/1000,cad.C_L,'coroa — só no módulo arara'),
 ('Tampa','PP · molde existe','—',cad.M_T/1000,cad.C_T,'uma por poste'),
]
KR={}
for i,(a,b,c,kg,cu,pap) in enumerate(pecas):
    r=5+i; KR[a]=r
    K.cell(r,1,a).font=fB; K.cell(r,2,b).font=fb; K.cell(r,3,c).font=fb
    x=K.cell(r,4,round(kg,4)); x.font=fin; x.number_format='0.0000'
    y=K.cell(r,5,round(cu,4)); y.font=fin; y.number_format='"R$" 0.0000'; y.fill=YEL
    K.cell(r,6,pap).font=fn
    for k in range(1,7): K.cell(r,k).border=box
r=5+len(pecas)+1
K.cell(r,1,'Saem: painel de 200 (56,8% da curva, e é o que empena), painéis 200×634/754, 300×360, 460×360 e 460×450 (a regra da proporção 1,3–2,6), ripas PSC-01 (315) e PSC-03 (595), porta-haste.').font=fn
K.cell(r,1).alignment=WR; K.merge_cells(start_row=r,start_column=1,end_row=r+1,end_column=6)
widths(K,[22,18,12,15,15,56])
def kd(n): return f'Kit!$D${KR[n]}'
def kc(n): return f'Kit!$E${KR[n]}'

# ============================================================ Modulos (45)
M=wb.create_sheet('Modulos')
M['A1']='Os 45 módulos padrão — 3 footprints × 3 larguras (N vãos) × 5 pilhas'; M['A1'].font=ft
M['A2']='N = 1 é módulo autônomo (quatro postes próprios, zero cruzeta, dois adultos carregam). N = 2 e 3 dividem poste (cruzeta) e são para parede fixa. Filtre pela linha 4.'; M['A2'].font=fn
cols=['ID','Painel','Ripa largura','Ripa comprimento','N vãos','Pilha','n 270','n 513','Coroa (mm)','Prateleiras',
      'Comprimento (mm)','Comprimento c/ coroa (mm)','Profundidade (mm)','Altura (mm)','Face da última prat. (mm)',
      'Trizetas','Cruzetas','Peças L','Tampas','Painéis','Ripas compr.','Ripas larg.','Ripas 270','Ripas 513','Pés',
      'Massa (kg)','Custo material (R$)','Frente linear (m)','Altura ÷ prof.','Uso']
header(M,4,cols)
FOOT=[('300×450',287,415,'Painel 300×450','Ripa largura 287','Ripa comprimento 415'),
      ('300×754',287,717,'Painel 300×754','Ripa largura 287','Ripa comprimento 717'),
      ('460×754',415,717,'Painel 460×754','Ripa largura 415','Ripa comprimento 717')]
PIL=[('baixo · 270×3',3,0,0,'ilha, checkout, mesa de demonstração — vê-se por cima'),
     ('médio · 270×5',5,0,0,'gôndola de miolo, ponta'),
     ('alto · 270×6',6,0,0,'parede: duas prateleiras na zona dos olhos'),
     ('alto-fundo · 513·513·270·270',2,2,0,'parede de banho/lavanderia: vão de 490 embaixo para lixeira e cesto'),
     ('arara · 513·513 + coroa 513',0,2,513,'arara e gancheira: a peça L fecha o topo sem prateleira')]
r=5; ids=[]
for fp,lw,cw,pan,rl,rc in FOOT:
    for N in (1,2,3):
        for nome,n270,n513,cor,uso in PIL:
            mid=f'{fp}-{N}-{nome.split(" ")[0]}'; ids.append(mid)
            M.cell(r,1,mid).font=fB
            M.cell(r,2,fp).font=fb; M.cell(r,3,lw).font=fin; M.cell(r,4,cw).font=fin; M.cell(r,5,N).font=fin
            M.cell(r,6,nome).font=fb; M.cell(r,7,n270).font=fin; M.cell(r,8,n513).font=fin; M.cell(r,9,cor).font=fin
            M.cell(r,10,f'=G{r}+H{r}+1')
            M.cell(r,11,f'=2*{N_["NOX"]}+(E{r}-1)*{N_["NOXC"]}+E{r}*(D{r}-{N_["CONS"]})').number_format=MM
            M.cell(r,12,f'=IF(I{r}>0,K{r}+2*{N_["NOL"]},K{r})').number_format=MM
            M.cell(r,13,f'=C{r}+2*({N_["NOY"]}-{N_["ENC"]})').number_format=MM
            M.cell(r,14,f'={N_["PE"]}+(J{r}+IF(I{r}>0,1,0))*{N_["NOZ"]}+G{r}*(270-{N_["CONS"]})+H{r}*(513-{N_["CONS"]})+IF(I{r}>0,I{r}-{N_["CONS"]},0)').number_format=MM
            M.cell(r,15,f'={N_["PE"]}+J{r}*{N_["NOZ"]}+G{r}*(270-{N_["CONS"]})+H{r}*(513-{N_["CONS"]})+{N_["PANT"]}/2').number_format=MM
            M.cell(r,16,f'=4*J{r}')                                    # trizetas
            M.cell(r,17,f'=2*(E{r}-1)*(J{r}+IF(I{r}>0,1,0))')          # cruzetas
            M.cell(r,18,f'=IF(I{r}>0,4,0)')                            # pecas L
            M.cell(r,19,f'=2*(E{r}+1)')                                # tampas
            M.cell(r,20,f'=E{r}*J{r}')                                 # paineis
            M.cell(r,21,f'=2*E{r}*J{r}+IF(I{r}>0,2*E{r},0)')          # ripas comprimento
            M.cell(r,22,f'=(E{r}+1)*J{r}')                             # ripas largura
            M.cell(r,23,f'=2*(E{r}+1)*(G{r}+IF(I{r}=270,1,0))')        # ripas 270
            M.cell(r,24,f'=2*(E{r}+1)*(H{r}+IF(I{r}=513,1,0))')        # ripas 513
            M.cell(r,25,f'=2*(E{r}+1)')                                # pes
            massa=(f'=P{r}*{kd("Trizeta")}+Q{r}*{kd("Cruzeta")}+R{r}*{kd("Peça L")}+S{r}*{kd("Tampa")}+T{r}*{kd(pan)}'
                   f'+U{r}*{kd(rc)}+V{r}*{kd(rl)}+W{r}*{kd("Ripa vertical 270")}+X{r}*{kd("Ripa vertical 513")}+Y{r}*{kd("Pé")}')
            M.cell(r,26,massa).number_format=KG
            M.cell(r,27,massa.replace('$D$','$E$')).number_format=BRL
            M.cell(r,28,f'=K{r}*J{r}/1000').number_format='0.00'
            M.cell(r,29,f'=N{r}/M{r}').number_format='0.00'
            M.cell(r,30,uso).font=fn
            for k in range(1,31):
                M.cell(r,k).border=box
                if k in (10,11,12,13,14,15,16,17,18,19,20,21,22,23,24,25,26,27,28,29): M.cell(r,k).alignment=C
            if N==1: 
                for k in range(1,31): M.cell(r,k).fill=GREY
            r+=1
LAST=r-1
M.cell(r+1,1,'Linhas cinza (N = 1): módulos autônomos, sem cruzeta — os que o showroom pode montar hoje. Arara não conta prateleira: os 3 níveis são a base e a coroa. Custo é só material; sem corte, montagem, embalagem, frete. Módulo contra parede não leva fundo.').font=fn
M.cell(r+1,1).alignment=WR; M.merge_cells(start_row=r+1,start_column=1,end_row=r+2,end_column=14)
widths(M,[26,9,8,9,6,28,6,6,7,8,12,12,12,10,12,8,8,7,7,8,8,8,8,8,6,9,13,9,8,52])
M.freeze_panes='B5'; M.auto_filter.ref=f'A4:AD{LAST}'

# ============================================================ Showroom
S=wb.create_sheet('Showroom')
S['A1']='Montagem sugerida do showroom só com módulos N = 1 (zero cruzeta)'; S['A1'].font=ft
S['A2']='Escolha o ID na coluna B (lista da aba Modulos) e a quantidade; o resto é fórmula. Quando a cruzeta existir, troque paredes fixas por N = 3 e a conta se refaz.'; S['A2'].font=fn
header(S,4,['Onde','ID do módulo','Qtd','Comprimento montado (mm)','Corrida livre (mm)','Folga (mm)','Custo (R$)','Frente (m)','Massa (kg)','Cruzetas'])
plano=[('Paredão sul · cozinha','300×754-1-alto',17,13350),
       ('Paredão norte · organização','300×754-1-alto',8,6100),
       ('Paredão do fundo · banho e lavanderia','460×754-1-alto-fundo',9,6873),
       ('Parede de entrada · frasqueiras','300×754-1-alto-fundo',8,6548),
       ('Folga caixa–pilar','300×450-1-alto',1,700),
       ('Ilha 1 · 2 × 2 costa a costa','460×754-1-baixo',4,None),
       ('Ilha 2 · 2 × 2 costa a costa','460×754-1-baixo',4,None),
       ('Corredor de checkout · 2 lados × 3','300×754-1-baixo',6,None),
       ('Arara Nitron-Mob no miolo','300×754-1-arara',2,None)]
LK=lambda col,r: f'INDEX(Modulos!${col}$5:${col}${LAST},MATCH($B{r},Modulos!$A$5:$A${LAST},0))'
for i,(onde,mid,q,livre) in enumerate(plano):
    r=5+i
    S.cell(r,1,onde).font=fB; S.cell(r,2,mid).font=fin; S.cell(r,3,q).font=fin
    S.cell(r,4,f'=C{r}*{LK("K",r)}').number_format=MM
    if livre: S.cell(r,5,livre).font=fin; S.cell(r,6,f'=E{r}-D{r}').number_format=MM
    S.cell(r,7,f'=C{r}*{LK("AA",r)}').number_format=BRL
    S.cell(r,8,f'=C{r}*{LK("AB",r)}').number_format='0.0'
    S.cell(r,9,f'=C{r}*{LK("Z",r)}').number_format=KG
    S.cell(r,10,f'=C{r}*{LK("Q",r)}')
    for k in range(1,11): S.cell(r,k).border=box
r=5+len(plano)
S.cell(r,1,'TOTAL').font=fB; S.cell(r,3,f'=SUM(C5:C{r-1})').font=fB
for c,fm in ((7,BRL),(8,'0.0'),(9,KG),(10,MM)):
    x=S.cell(r,c,f'=SUM({L(c)}5:{L(c)}{r-1})'); x.font=fB; x.number_format=fm
for k in range(1,11): S.cell(r,k).border=box
S.cell(r+2,1,'Peças a comprar / injetar').font=fB
header(S,r+3,['Peça','Quantidade','Massa (kg)','Custo (R$)'])
comp=[('Trizeta','P'),('Cruzeta','Q'),('Peça L','R'),('Tampa','S'),('Painel 300×450',None),('Painel 300×754',None),('Painel 460×754',None),
      ('Ripa largura 287',None),('Ripa largura 415',None),('Ripa comprimento 415',None),('Ripa comprimento 717',None),('Ripa vertical 270','W'),('Ripa vertical 513','X'),('Pé','Y')]
for i,(nome,col) in enumerate(comp):
    rr=r+4+i; S.cell(rr,1,nome).font=fB
    if col:
        S.cell(rr,2,'=SUMPRODUCT($C$5:$C$%d,'%(r-1)+','.join([])+f'IFERROR(INDEX(Modulos!${col}$5:${col}${LAST},MATCH($B$5:$B${r-1},Modulos!$A$5:$A${LAST},0)),0))')
    elif nome.startswith('Painel'):
        fp=nome.split(' ')[1]
        S.cell(rr,2,f'=SUMPRODUCT($C$5:$C${r-1},--(LEFT($B$5:$B${r-1},7)="{fp}"),IFERROR(INDEX(Modulos!$T$5:$T${LAST},MATCH($B$5:$B${r-1},Modulos!$A$5:$A${LAST},0)),0))')
    elif nome=='Ripa largura 287':
        S.cell(rr,2,f'=SUMPRODUCT($C$5:$C${r-1},--(LEFT($B$5:$B${r-1},3)="300"),IFERROR(INDEX(Modulos!$V$5:$V${LAST},MATCH($B$5:$B${r-1},Modulos!$A$5:$A${LAST},0)),0))')
    elif nome=='Ripa largura 415':
        S.cell(rr,2,f'=SUMPRODUCT($C$5:$C${r-1},--(LEFT($B$5:$B${r-1},3)="460"),IFERROR(INDEX(Modulos!$V$5:$V${LAST},MATCH($B$5:$B${r-1},Modulos!$A$5:$A${LAST},0)),0))')
    elif nome=='Ripa comprimento 415':
        S.cell(rr,2,f'=SUMPRODUCT($C$5:$C${r-1},--(MID($B$5:$B${r-1},5,3)="450"),IFERROR(INDEX(Modulos!$U$5:$U${LAST},MATCH($B$5:$B${r-1},Modulos!$A$5:$A${LAST},0)),0))')
    elif nome=='Ripa comprimento 717':
        S.cell(rr,2,f'=SUMPRODUCT($C$5:$C${r-1},--(MID($B$5:$B${r-1},5,3)="754"),IFERROR(INDEX(Modulos!$U$5:$U${LAST},MATCH($B$5:$B${r-1},Modulos!$A$5:$A${LAST},0)),0))')
    S.cell(rr,2).number_format=MM
    S.cell(rr,3,f'=B{rr}*{kd(nome)}').number_format=KG
    S.cell(rr,4,f'=B{rr}*{kc(nome)}').number_format=BRL
    for k in range(1,5): S.cell(rr,k).border=box
widths(S,[38,24,6,22,16,10,14,10,10,9])
wb.move_sheet('Parametros',offset=3)
wb.save(OUT); print('gravado',OUT,'·',len(ids),'módulos')
